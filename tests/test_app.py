"""
Integration tests for physdb Flask app.

Coverage areas:
  - Authentication: login success/failure, logout
  - Open redirect rejection in login `next` parameter
  - enforce_password_change blocks all routes except allowed endpoints
  - Role enforcement: viewer vs physicist on manage_equipment_required routes
  - CSV row limit rejection (>500 rows)
  - get_estimated_eol_date date arithmetic
  - Admin database export (.sqlite3 backup and .xlsx workbook)
  - Full Facility: migration, admin edit, details page, import/export
  - Equipment import: export round trip, partial updates, text parsing, per-row errors
  - Security: personnel permission limits, CSRF enforcement, output escaping
  - Shared equipment filters, parameter hardening, inline edits, capital category overlap
"""
import csv
import io
import os
import re
import sqlite3
import tempfile
import openpyxl
import pytest
from datetime import date
from unittest.mock import MagicMock, patch

from flask import g
from sqlalchemy import inspect, text

from app import db, check_and_migrate_db, login_throttle, ComplianceTest, Department, Equipment, EquipmentClass, Facility, Personnel
from conftest import make_user, login


# ---------------------------------------------------------------------------
# Authentication
# ---------------------------------------------------------------------------

class TestLogin:
    def test_valid_login_redirects_to_index(self, app, client):
        with app.app_context():
            make_user(username='alice', password='correct')
        resp = login(client, 'alice', 'correct')
        assert resp.status_code == 302
        assert resp.location.endswith('/')

    def test_wrong_password_stays_on_login(self, app, client):
        with app.app_context():
            make_user(username='bob', password='correct')
        resp = login(client, 'bob', 'wrong')
        assert resp.status_code == 200
        assert b'Invalid username or password' in resp.data

    def test_inactive_user_cannot_login(self, app, client):
        with app.app_context():
            make_user(username='carol', password='pw', is_active=False)
        resp = login(client, 'carol', 'pw')
        assert resp.status_code == 200
        assert b'Invalid username or password' in resp.data

    def test_login_required_false_cannot_login(self, app, client):
        """Personnel with login_required=False should not be able to log in."""
        with app.app_context():
            make_user(username='dave', password='pw', login_required=False)
        resp = login(client, 'dave', 'pw')
        assert resp.status_code == 200
        assert b'Invalid username or password' in resp.data


# ---------------------------------------------------------------------------
# Open redirect
# ---------------------------------------------------------------------------

class TestOpenRedirect:
    def test_external_next_param_ignored(self, app, client):
        with app.app_context():
            make_user(username='eve', password='pw')
        resp = client.post(
            '/login?next=https://evil.example.com/steal',
            data={'username': 'eve', 'password': 'pw'},
            follow_redirects=False,
        )
        assert resp.status_code == 302
        location = resp.location
        assert 'evil.example.com' not in location

    def test_relative_next_param_accepted(self, app, client):
        with app.app_context():
            make_user(username='frank', password='pw')
        resp = client.post(
            '/login?next=/equipment',
            data={'username': 'frank', 'password': 'pw'},
            follow_redirects=False,
        )
        assert resp.status_code == 302
        assert 'evil' not in resp.location


# ---------------------------------------------------------------------------
# Unauthenticated access
# ---------------------------------------------------------------------------

class TestUnauthenticatedAccess:
    PROTECTED_ROUTES = [
        '/equipment/new',
        '/import-data',
        '/import-personnel',
        '/export-equipment',
        '/change-password',
    ]

    @pytest.mark.parametrize('route', PROTECTED_ROUTES)
    def test_redirects_to_login(self, client, route):
        resp = client.get(route, follow_redirects=False)
        assert resp.status_code == 302
        assert '/login' in resp.location


# ---------------------------------------------------------------------------
# enforce_password_change
# ---------------------------------------------------------------------------

class TestPasswordChangeEnforcement:
    def _login_must_change(self, app, client):
        with app.app_context():
            make_user(username='grace', password='pw', must_change_password=True)
        login(client, 'grace', 'pw')

    def test_blocked_from_index(self, app, client):
        self._login_must_change(app, client)
        resp = client.get('/', follow_redirects=False)
        assert resp.status_code == 302
        assert 'change-password' in resp.location

    def test_change_password_itself_is_allowed(self, app, client):
        self._login_must_change(app, client)
        resp = client.get('/change-password', follow_redirects=False)
        # Should render the form, not redirect away
        assert resp.status_code == 200

    def test_logout_is_allowed(self, app, client):
        self._login_must_change(app, client)
        resp = client.post('/logout', follow_redirects=False)
        assert resp.status_code == 302
        assert '/login' in resp.location


class TestLoginHardening:
    def test_logout_requires_post(self, app, client):
        with app.app_context():
            make_user(username='alice', password='pw')
        login(client, 'alice', 'pw')
        assert client.get('/logout').status_code == 405
        assert client.get('/', follow_redirects=False).status_code == 200  # still logged in

    def test_repeated_failures_are_throttled(self, app, client):
        with app.app_context():
            make_user(username='alice', password='correct')
        for _ in range(login_throttle.max_failures):
            assert login(client, 'alice', 'wrong').status_code == 200
        # Even the right password is refused while the IP is blocked
        resp = login(client, 'alice', 'correct')
        assert resp.status_code == 429
        assert b'Too many failed login attempts' in resp.data

    def test_success_clears_failures(self, app, client):
        with app.app_context():
            make_user(username='alice', password='correct')
        for _ in range(login_throttle.max_failures - 1):
            login(client, 'alice', 'wrong')
        assert login(client, 'alice', 'correct').status_code == 302
        client.post('/logout')
        assert login(client, 'alice', 'wrong').status_code == 200  # counter restarted, not blocked


# ---------------------------------------------------------------------------
# Role enforcement
# ---------------------------------------------------------------------------

class TestRoleEnforcement:
    def test_viewer_cannot_access_equipment_new(self, app, client):
        """A user with no manage roles should be redirected away from /equipment/new."""
        with app.app_context():
            make_user(username='viewer', password='pw', roles='viewer')
        login(client, 'viewer', 'pw')
        resp = client.get('/equipment/new', follow_redirects=False)
        assert resp.status_code == 302
        assert 'equipment' in resp.location  # redirected to equipment_list

    def test_physicist_can_access_equipment_new(self, app, client):
        with app.app_context():
            make_user(username='phys', password='pw', roles='physicist')
        login(client, 'phys', 'pw')
        resp = client.get('/equipment/new', follow_redirects=False)
        # Should render the form (200), not be redirected
        assert resp.status_code == 200

    def test_admin_can_access_equipment_new(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
        login(client, 'admin', 'pw')
        resp = client.get('/equipment/new', follow_redirects=False)
        assert resp.status_code == 200

    def test_viewer_cannot_access_import_personnel(self, app, client):
        with app.app_context():
            make_user(username='viewer2', password='pw', roles='viewer')
        login(client, 'viewer2', 'pw')
        resp = client.get('/import-personnel', follow_redirects=False)
        assert resp.status_code == 302


# ---------------------------------------------------------------------------
# CSV row limit
# ---------------------------------------------------------------------------

def _make_csv(num_rows, columns=None):
    """Return a BytesIO CSV with the given number of data rows."""
    if columns is None:
        columns = ['equipment_class', 'eq_mod']
    buf = io.BytesIO()
    header = ','.join(columns) + '\n'
    buf.write(header.encode())
    for i in range(num_rows):
        row = ','.join([f'value{i}'] * len(columns)) + '\n'
        buf.write(row.encode())
    buf.seek(0)
    return buf


class TestCSVRowLimit:
    def _login_physicist(self, app, client):
        with app.app_context():
            make_user(username='phys_csv', password='pw', roles='physicist')
        login(client, 'phys_csv', 'pw')

    def test_import_data_rejects_over_500_rows(self, app, client):
        self._login_physicist(app, client)
        csv_data = _make_csv(501)
        resp = client.post(
            '/import-data',
            data={'file': (csv_data, 'big.csv', 'text/csv')},
            content_type='multipart/form-data',
            follow_redirects=True,
        )
        assert b'exceeds the maximum of 500 rows' in resp.data

    def test_import_data_accepts_500_rows(self, app, client):
        self._login_physicist(app, client)
        csv_data = _make_csv(500)
        resp = client.post(
            '/import-data',
            data={'file': (csv_data, 'ok.csv', 'text/csv')},
            content_type='multipart/form-data',
            follow_redirects=True,
        )
        assert b'exceeds the maximum' not in resp.data

    def test_import_personnel_rejects_over_500_rows(self, app, client):
        with app.app_context():
            make_user(username='phys_pcsv', password='pw', roles='physicist')
        login(client, 'phys_pcsv', 'pw')
        csv_data = _make_csv(501, columns=['name', 'email'])
        resp = client.post(
            '/import-personnel',
            data={'csv_file': (csv_data, 'big.csv', 'text/csv')},
            content_type='multipart/form-data',
            follow_redirects=True,
        )
        assert b'exceeds the maximum of 500 rows' in resp.data


# ---------------------------------------------------------------------------
# Date arithmetic — get_estimated_eol_date
# ---------------------------------------------------------------------------

class TestEstimatedEolDate:
    """
    get_estimated_eol_date picks the latest of eq_mandt/eq_instdt/eq_rfrbdt
    and adds equipment_subclass.expected_lifetime years.
    """

    def _make_equipment_stub(self, mandt=None, instdt=None, rfrbdt=None, lifetime=None):
        """A plain object with just the attributes get_estimated_eol_date reads.

        A real Equipment built with __new__ has no SQLAlchemy instance state, so
        setting its columns fails; binding the method to a stub avoids that.
        """
        from types import SimpleNamespace
        from app import Equipment
        stub = SimpleNamespace(eq_mandt=mandt, eq_instdt=instdt, eq_rfrbdt=rfrbdt,
                               equipment_subclass=SimpleNamespace(expected_lifetime=lifetime))
        stub.get_estimated_eol_date = Equipment.get_estimated_eol_date.__get__(stub)
        return stub

    def test_returns_none_without_subclass(self, app):
        with app.app_context():
            eq = self._make_equipment_stub(mandt=date(2010, 1, 1))
            eq.equipment_subclass = None
            assert eq.get_estimated_eol_date() is None

    def test_returns_none_without_dates(self, app):
        with app.app_context():
            eq = self._make_equipment_stub(lifetime=10)
            assert eq.get_estimated_eol_date() is None

    def test_uses_manufacture_date(self, app):
        with app.app_context():
            eq = self._make_equipment_stub(mandt=date(2010, 6, 15), lifetime=10)
            result = eq.get_estimated_eol_date()
            assert result == date(2020, 6, 15)

    def test_uses_latest_of_mandt_and_instdt(self, app):
        with app.app_context():
            # instdt is later — should be used as the base
            eq = self._make_equipment_stub(
                mandt=date(2010, 1, 1),
                instdt=date(2012, 3, 20),
                lifetime=5,
            )
            result = eq.get_estimated_eol_date()
            assert result == date(2017, 3, 20)

    def test_rfrbdt_takes_precedence_when_latest(self, app):
        with app.app_context():
            eq = self._make_equipment_stub(
                mandt=date(2010, 1, 1),
                instdt=date(2012, 1, 1),
                rfrbdt=date(2015, 6, 1),
                lifetime=8,
            )
            result = eq.get_estimated_eol_date()
            assert result == date(2023, 6, 1)

    def test_lifetime_zero_means_no_estimate(self, app):
        # A 0-year lifetime is treated as "not set" rather than making the unit due on day one
        with app.app_context():
            eq = self._make_equipment_stub(mandt=date(2015, 1, 1), lifetime=0)
            assert eq.get_estimated_eol_date() is None


# ---------------------------------------------------------------------------
# Admin database export
# ---------------------------------------------------------------------------

class TestDatabaseExport:
    """The /admin/database/* endpoints: access control, file validity, secrets."""

    BACKUP_URL = '/admin/database/backup'
    WORKBOOK_URL = '/admin/database/workbook'

    @pytest.mark.parametrize('url', [BACKUP_URL, WORKBOOK_URL])
    def test_anonymous_is_redirected_to_login(self, app, client, url):
        resp = client.get(url)
        assert resp.status_code == 302
        assert '/login' in resp.location

    @pytest.mark.parametrize('url', [BACKUP_URL, WORKBOOK_URL])
    def test_non_admin_is_rejected(self, app, client, url):
        with app.app_context():
            make_user(username='viewer', password='pw', is_admin=False)
        login(client, 'viewer', 'pw')
        resp = client.get(url, follow_redirects=True)
        assert b'Admin access required.' in resp.data

    def test_backup_returns_a_valid_sqlite_file(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
        login(client, 'admin', 'pw')

        resp = client.get(self.BACKUP_URL)
        assert resp.status_code == 200
        assert 'attachment' in resp.headers['Content-Disposition']
        assert '.sqlite3' in resp.headers['Content-Disposition']
        assert resp.data.startswith(b'SQLite format 3\x00')

        # The snapshot must actually open and contain the schema.
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'restored.sqlite3')
            with open(path, 'wb') as fh:
                fh.write(resp.data)
            conn = sqlite3.connect(path)
            try:
                names = {r[0] for r in conn.execute(
                    "SELECT name FROM sqlite_master WHERE type='table'")}
            finally:
                conn.close()
        assert {'equipment', 'personnel', 'compliance_tests'} <= names

    def test_backup_retains_password_hashes(self, app, client):
        """A backup you cannot restore logins from is not a backup."""
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
        login(client, 'admin', 'pw')

        resp = client.get(self.BACKUP_URL)
        assert b'password_hash' in resp.data

    def test_workbook_has_one_sheet_per_table(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
        login(client, 'admin', 'pw')

        resp = client.get(self.WORKBOOK_URL)
        assert resp.status_code == 200
        assert resp.headers['Content-Type'] == (
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        assert '.xlsx' in resp.headers['Content-Disposition']
        assert resp.data.startswith(b'PK')  # xlsx is a zip container

        workbook = openpyxl.load_workbook(io.BytesIO(resp.data), read_only=True)
        try:
            sheets = workbook.sheetnames
        finally:
            workbook.close()
        assert sheets[0] == '_README'
        assert {'equipment', 'personnel', 'equipment_classes'} <= set(sheets)

    def test_workbook_omits_password_hash(self, app, client):
        """The one regression that actually matters."""
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
        login(client, 'admin', 'pw')

        resp = client.get(self.WORKBOOK_URL)
        workbook = openpyxl.load_workbook(io.BytesIO(resp.data), read_only=True)
        try:
            headers = [c.value for c in next(workbook['personnel'].iter_rows(max_row=1))]
        finally:
            workbook.close()
        assert 'username' in headers
        assert 'password_hash' not in headers

    def test_workbook_keeps_foreign_keys_as_integers(self, app, client):
        """pandas would coerce a nullable int column to float ('3.0'); we must not."""
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            cls = EquipmentClass(name='CT')
            db.session.add(cls)
            db.session.commit()
            db.session.add(Equipment(class_id=cls.id, eq_mod='Optima'))
            db.session.commit()
        login(client, 'admin', 'pw')

        resp = client.get(self.WORKBOOK_URL)
        workbook = openpyxl.load_workbook(io.BytesIO(resp.data), read_only=True)
        try:
            rows = list(workbook['equipment'].iter_rows(max_row=2, values_only=True))
        finally:
            workbook.close()
        class_id = rows[1][rows[0].index('class_id')]
        assert isinstance(class_id, int) and not isinstance(class_id, bool)

    def test_backup_unavailable_on_non_sqlite_backend(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
        login(client, 'admin', 'pw')

        with patch('app._is_sqlite_backend', return_value=False):
            resp = client.get(self.BACKUP_URL, follow_redirects=True)
        assert b'only available for SQLite' in resp.data


# ---------------------------------------------------------------------------
# Full Facility (facilities.facility_full)
# ---------------------------------------------------------------------------

class TestLegacyMigration:
    """check_and_migrate_db() cleans up data from older versions on startup."""

    def test_old_lowercase_annual_test_still_drives_next_due_date(self, app):
        cls = EquipmentClass(name='CT')
        db.session.add(cls)
        db.session.commit()
        eq = Equipment(class_id=cls.id, eq_auditfreq='Annual - TJC')
        db.session.add(eq)
        db.session.commit()
        # Stored the way old versions did, bypassing the form
        with db.engine.begin() as conn:
            conn.execute(text("INSERT INTO compliance_tests (eq_id, test_type, test_date) "
                              "VALUES (:eq, 'annual', '2024-03-01')"), {'eq': eq.eq_id})
        check_and_migrate_db()
        db.session.expire_all()
        assert ComplianceTest.query.one().test_type == 'Annual'
        assert eq.get_last_tested_date() == date(2024, 3, 1)
        assert eq.get_next_due_date() == date(2025, 3, 31)  # 1 year + 30 days

    def test_unused_equipment_columns_are_dropped(self, app):
        with db.engine.begin() as conn:
            for col, kind in (('eq_address', 'TEXT'), ('eq_eeoldate', 'DATE'),
                              ('eq_capcat', 'INTEGER'), ('eq_capecst', 'INTEGER')):
                conn.execute(text(f'ALTER TABLE equipment ADD COLUMN {col} {kind}'))
        check_and_migrate_db()
        check_and_migrate_db()
        cols = {c['name'] for c in inspect(db.engine).get_columns('equipment')}
        assert not cols & {'eq_address', 'eq_eeoldate', 'eq_capcat', 'eq_capecst'}

    def test_admin_role_is_removed_but_admin_rights_kept(self, app):
        admin = make_user(username='admin', password='pw', is_admin=True, roles='admin, physicist')
        check_and_migrate_db()
        db.session.expire_all()
        assert admin.get_roles_list() == ['physicist']
        assert admin.is_admin is True


class TestFacilityFull:
    def _make_equipment(self, facility_full=None):
        cls = EquipmentClass(name='CT')
        fac = Facility(name='St. Mary', facility_full=facility_full)
        db.session.add_all([cls, fac])
        db.session.commit()
        eq = Equipment(class_id=cls.id, facility_id=fac.id)
        db.session.add(eq)
        db.session.commit()
        return eq.eq_id

    def test_migration_adds_column_and_is_idempotent(self, app):
        with db.engine.begin() as conn:
            conn.execute(text('ALTER TABLE facilities DROP COLUMN facility_full'))
        check_and_migrate_db()
        check_and_migrate_db()
        cols = {c['name'] for c in inspect(db.engine).get_columns('facilities')}
        assert 'facility_full' in cols

    def test_admin_edit_saves_facility_full(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            fac = Facility(name='St. Mary')
            db.session.add(fac)
            db.session.commit()
            fac_id = fac.id
        login(client, 'admin', 'pw')

        client.post(f'/admin/facilities/{fac_id}/edit',
                    data={'name': 'St. Mary', 'facility_full': "St. Mary's Regional Medical Center"})
        assert db.session.get(Facility, fac_id).facility_full == "St. Mary's Regional Medical Center"

    def test_admin_rejects_duplicate_facility_full(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            db.session.add(Facility(name='A', facility_full='Shared Full Name'))
            fac_b = Facility(name='B')
            db.session.add(fac_b)
            db.session.commit()
            fac_b_id = fac_b.id
        login(client, 'admin', 'pw')

        resp = client.post(f'/admin/facilities/{fac_b_id}/edit',
                           data={'name': 'B', 'facility_full': 'Shared Full Name'}, follow_redirects=True)
        assert b'already used by another facility' in resp.data
        assert db.session.get(Facility, fac_b_id).facility_full is None

    def test_detail_page_shows_facility_full(self, app, client):
        with app.app_context():
            make_user(username='viewer', password='pw')
            eq_id = self._make_equipment("St. Mary's Regional Medical Center")
        login(client, 'viewer', 'pw')

        resp = client.get(f'/equipment/{eq_id}')
        assert b'Full Facility:' in resp.data
        assert b'St. Mary&#39;s Regional Medical Center' in resp.data

    def test_detail_page_hides_blank_facility_full(self, app, client):
        with app.app_context():
            make_user(username='viewer', password='pw')
            eq_id = self._make_equipment()
        login(client, 'viewer', 'pw')

        resp = client.get(f'/equipment/{eq_id}')
        assert b'Full Facility:' not in resp.data

    def test_equipment_export_includes_facility_full(self, app, client):
        with app.app_context():
            make_user(username='viewer', password='pw')
            self._make_equipment('Full Name Here')
        login(client, 'viewer', 'pw')

        resp = client.get('/export-equipment')
        header, first = resp.data.decode().splitlines()[:2]
        cols = header.split(',')
        assert first.split(',')[cols.index('facility_full')] == 'Full Name Here'

    def test_facility_import_without_column_keeps_existing_value(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            fac = Facility(name='St. Mary', facility_full='Keep Me')
            db.session.add(fac)
            db.session.commit()
            fac_id = fac.id
        login(client, 'admin', 'pw')

        csv_file = (io.BytesIO(b'name,address\nSt. Mary,1 Main St\n'), 'facilities.csv')
        client.post('/import-facilities', data={'csv_file': csv_file},
                    content_type='multipart/form-data')
        assert db.session.get(Facility, fac_id).facility_full == 'Keep Me'


# ---------------------------------------------------------------------------
# Equipment import: export -> edit -> import round trip, and blank CSV cells
# ---------------------------------------------------------------------------

def _post_import(client, text, url='/import-data'):
    return client.post(url, data={'file': (io.BytesIO(text.encode('utf-8')), 'e.csv')},
                       content_type='multipart/form-data', follow_redirects=True)


class TestImportRoundTrip:
    def _login(self, app, client):
        with app.app_context():
            make_user(username='phys', password='pw', roles='physicist')
            cls = EquipmentClass(name='CT')
            db.session.add_all([cls, Facility(name='Old Site'), Facility(name='New Site')])
            db.session.commit()
            db.session.add(Equipment(class_id=cls.id, facility_id=1, eq_rm='12', eq_notes='keep',
                                     eq_rfrbdt=date(2019, 1, 1), eq_physcov=True))
            db.session.commit()
        login(client, 'phys', 'pw')

    def test_exported_csv_edits_apply(self, app, client):
        """The import reads every column /export-equipment writes, including ones bulk edit ignored."""
        self._login(app, client)
        rows = list(csv.DictReader(io.StringIO(client.get('/export-equipment').data.decode())))
        rows[0].update({'facility': 'New Site', 'equipment_subclass': 'Head', 'manufacturer': 'Acme',
                        'eq_rfrbdt': '2021-05-01', 'eq_physcov': 'FALSE', 'eq_capyr': '2030',
                        'eq_captype': 'Upgrade', 'eq_capnote': 'note'})
        out = io.StringIO()
        writer = csv.DictWriter(out, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)
        resp = _post_import(client, out.getvalue())
        assert b'updated 1 existing' in resp.data
        assert b'unrecognized' not in resp.data

        eq = db.session.get(Equipment, 1)
        assert eq.facility.name == 'New Site'
        assert eq.manufacturer.name == 'Acme'
        assert eq.equipment_subclass.name == 'Head'
        assert eq.equipment_subclass.class_id == eq.class_id
        assert eq.eq_rfrbdt == date(2021, 5, 1)
        assert eq.eq_physcov is False
        assert (eq.eq_capyr, eq.eq_captype, eq.eq_capnote) == (2030, 'Upgrade', 'note')
        assert eq.eq_notes == 'keep' and eq.eq_rm == '12'
        assert Equipment.query.count() == 1

    def test_missing_columns_leave_fields_alone_and_blank_cells_clear(self, app, client):
        self._login(app, client)
        _post_import(client, 'eq_id,eq_notes\n1,\n')
        eq = db.session.get(Equipment, 1)
        assert eq.eq_notes is None           # blank cell cleared it
        assert eq.eq_rm == '12'              # column absent: untouched
        assert eq.facility.name == 'Old Site'
        assert eq.eq_rfrbdt == date(2019, 1, 1)

    def test_text_is_not_turned_into_numbers(self, app, client):
        self._login(app, client)
        _post_import(client, 'equipment_class,eq_rm,eq_sn,eq_phone\nCT,101,00123,2076625555\nCT,,,\n')
        eq = Equipment.query.filter_by(eq_sn='00123').one()
        assert (eq.eq_rm, eq.eq_phone) == ('101', '2076625555')

    def test_excel_bom_and_dates(self, app, client):
        self._login(app, client)
        _post_import(client, '﻿eq_id,eq_instdt\n1,3/7/2022\n')
        assert db.session.get(Equipment, 1).eq_instdt == date(2022, 3, 7)

    def test_bad_row_is_reported_and_others_still_import(self, app, client):
        self._login(app, client)
        resp = _post_import(client, 'equipment_class,eq_mod\nCT,Good1\n,NoClass\nCT,Good2\n')
        assert b'Row 3: equipment_class is required' in resp.data
        assert sorted(e.eq_mod for e in Equipment.query.filter(Equipment.eq_mod.isnot(None))) == ['Good1', 'Good2']

    def test_new_contact_without_email_warns_instead_of_failing(self, app, client):
        self._login(app, client)
        resp = _post_import(client, 'equipment_class,eq_mod,contact_person\nCT,X,Jane Doe\n')
        assert b'needs an email' in resp.data
        eq = Equipment.query.filter_by(eq_mod='X').one()
        assert eq.contact_id is None

    def test_invalid_values_warn_and_leave_field_unchanged(self, app, client):
        self._login(app, client)
        resp = _post_import(client, 'eq_id,eq_instdt,eq_retired,eq_auditfreq,eq_capcst\n1,soon,maybe,Weekly,lots\n')
        for text in (b'eq_instdt', b'eq_retired', b'eq_auditfreq', b'eq_capcst'):
            assert text in resp.data
        eq = db.session.get(Equipment, 1)
        assert eq.eq_instdt is None and eq.eq_retired is False

    def test_unknown_columns_are_reported(self, app, client):
        self._login(app, client)
        resp = _post_import(client, 'eq_id,Facility,eq_notes\n1,New Site,x\n')
        assert b'Ignored unrecognized columns: Facility' in resp.data
        assert db.session.get(Equipment, 1).facility.name == 'Old Site'

    def test_subclass_is_matched_within_its_class(self, app, client):
        from app import EquipmentSubclass
        self._login(app, client)
        with app.app_context():
            mr = EquipmentClass(name='MR')
            db.session.add(mr)
            db.session.commit()
            db.session.add(EquipmentSubclass(name='Head', class_id=mr.id))
            db.session.commit()
        _post_import(client, 'eq_id,equipment_subclass\n1,Head\n')
        eq = db.session.get(Equipment, 1)
        assert eq.equipment_subclass.class_id == eq.class_id  # created under CT, not reused from MR

    def test_formula_text_is_escaped_on_export_and_restored_on_import(self, app, client):
        self._login(app, client)
        with app.app_context():
            db.session.get(Equipment, 1).eq_notes = '=1+2'
            db.session.commit()
        exported = client.get('/export-equipment').data.decode()
        assert "'=1+2" in exported
        _post_import(client, exported)
        assert db.session.get(Equipment, 1).eq_notes == '=1+2'


    def test_id_only_import_rebuilds_stale_mefacreg(self, app, client):
        # Older imports copied eq_mefacreg from the file (often a '-' placeholder);
        # re-importing just the IDs rebuilds it from eq_mefac and eq_mereg
        with app.app_context():
            make_user(username='p', password='pw', roles='physicist')
            db.session.add(EquipmentClass(name='US'))
            db.session.commit()
            with db.engine.begin() as conn:
                conn.execute(text("INSERT INTO equipment (class_id, eq_mod, eq_mefacreg) VALUES (1, 'US', '-')"))
                conn.execute(text("INSERT INTO equipment (class_id, eq_mod, eq_mefac, eq_mereg, eq_mefacreg) "
                                  "VALUES (1, 'CT', 'FAC-12', 'REG-345', 'old')"))
        login(client, 'p', 'pw')
        _post_import(client, 'eq_id\n1\n2\n')
        db.session.expire_all()
        assert [(e.eq_mod, e.eq_mefacreg) for e in Equipment.query.order_by(Equipment.eq_id)] == \
            [('US', None), ('CT', '12-345')]

class TestOtherImporters:
    def _setup(self, app, client):
        with app.app_context():
            admin = make_user(username='admin', password='pw', is_admin=True, roles='physicist')
            cls = EquipmentClass(name='CT')
            db.session.add(cls)
            db.session.commit()
            db.session.add(Equipment(class_id=cls.id))
            db.session.commit()
            admin_id = admin.id
        login(client, 'admin', 'pw')
        return admin_id

    def _upload(self, client, url, text, field='csv_file'):
        return client.post(url, data={field: (io.BytesIO(text.encode()), 'x.csv')},
                           content_type='multipart/form-data', follow_redirects=True)

    def test_compliance_blank_keeps_and_clear_empties(self, app, client):
        admin_id = self._setup(app, client)
        self._upload(client, '/import-compliance',
                     f'eq_id,test_type,test_date,report_date,notes,reviewed_by_id\n'
                     f'1,Annual,2024-01-01,2024-01-15,first,{admin_id}\n')
        test = ComplianceTest.query.one()
        resp = self._upload(client, '/import-compliance',
                            f'test_id,eq_id,test_type,test_date,report_date,notes,performed_by_id\n'
                            f'{test.test_id},1,Annual,1/2/2024,CLEAR,,999\n')
        assert b'updated 1 existing' in resp.data
        assert b'performed_by_id: no personnel' in resp.data
        test = db.session.get(ComplianceTest, test.test_id)
        assert test.test_date == date(2024, 1, 2)
        assert test.report_date is None          # CLEAR emptied it
        assert test.notes == 'first'             # blank left it alone
        assert test.reviewed_by_id == admin_id   # column absent: untouched

    def test_compliance_bad_rows_are_reported(self, app, client):
        self._setup(app, client)
        resp = self._upload(client, '/import-compliance',
                            'eq_id,test_type,test_date\n42,Annual,2024-01-01\n1,Annual,someday\n1,Annual,2024-03-01\n')
        assert b"Row 2: equipment ID &#39;42&#39; not found" in resp.data
        assert b'Row 3: test_date' in resp.data
        assert [t.test_date for t in ComplianceTest.query.all()] == [date(2024, 3, 1)]

    def test_compliance_export_round_trip_keeps_personnel(self, app, client):
        admin_id = self._setup(app, client)
        self._upload(client, '/import-compliance',
                     f'eq_id,test_type,test_date,performed_by_id,reviewed_by_id\n1,Annual,2024-01-01,{admin_id},{admin_id}\n')
        exported = client.get('/export-compliance').data.decode()
        row = next(csv.DictReader(io.StringIO(exported)))
        assert row['performed_by_id'] == str(admin_id) and row['reviewed_by'] == 'Test User'
        resp = self._upload(client, '/import-compliance', exported)
        assert b'updated 1 existing' in resp.data and b'unrecognized' not in resp.data
        test = ComplianceTest.query.one()
        assert (test.performed_by_id, test.reviewed_by_id) == (admin_id, admin_id)

    def test_scheduled_import_creates_and_updates(self, app, client):
        admin_id = self._setup(app, client)
        self._upload(client, '/import-scheduled-tests',
                     'eq_id,scheduled_date,scheduling_date,notes\n1,2030-01-01,2029-12-01,go\n', field='file')
        from app import ScheduledTest
        sched = ScheduledTest.query.one()
        assert (sched.created_by_id, sched.notes) == (admin_id, 'go')
        self._upload(client, '/import-scheduled-tests',
                     f'schedule_id,eq_id,scheduled_date,scheduling_date\n{sched.schedule_id},1,2031-01-01,2029-12-01\n',
                     field='file')
        assert ScheduledTest.query.one().scheduled_date == date(2031, 1, 1)

    def test_facility_import_without_is_active_keeps_status(self, app, client):
        self._setup(app, client)
        with app.app_context():
            db.session.add(Facility(name='Closed', is_active=False))
            db.session.commit()
        self._upload(client, '/import-facilities', 'name,address\nClosed,1 Main St\n')
        fac = Facility.query.filter_by(name='Closed').one()
        assert fac.is_active is False and fac.address == '1 Main St'

    def test_personnel_export_round_trip_keeps_roles(self, app, client):
        self._setup(app, client)
        with app.app_context():
            make_user(username='pa', password='pw', roles='physics_assistant, qa_technologist')
        exported = client.get('/export-personnel').data.decode()
        resp = self._upload(client, '/import-personnel', exported)
        assert b'unrecognized' not in resp.data
        person = Personnel.query.filter_by(username='pa').one()
        assert sorted(person.get_roles_list()) == ['physics_assistant', 'qa_technologist']

    def test_compliance_import_rejects_unknown_test_type(self, app, client):
        with app.app_context():
            make_user(username='phys', password='pw', roles='physicist')
            cls = EquipmentClass(name='CT')
            db.session.add(cls)
            db.session.commit()
            db.session.add(Equipment(class_id=cls.id))
            db.session.commit()
        login(client, 'phys', 'pw')
        csv_file = io.BytesIO(b"eq_id,test_type,test_date\n1,Annual,2024-01-01\n1,<script>,2024-02-01\n")
        client.post('/import-compliance', data={'csv_file': (csv_file, 'c.csv')},
                    content_type='multipart/form-data')
        assert [t.test_type for t in ComplianceTest.query.all()] == ['Annual']

    def test_personnel_phone_is_kept_as_text(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
        login(client, 'admin', 'pw')
        csv_file = io.BytesIO(b'name,email,phone,contact\nNew Person,new@example.com,2076625555,TRUE\n')
        client.post('/import-personnel', data={'csv_file': (csv_file, 'p.csv')},
                    content_type='multipart/form-data')
        assert Personnel.query.filter_by(email='new@example.com').one().phone == '2076625555'


class TestBlankCsvCells:
    def test_import_blank_facility_address_is_empty_not_nan(self, app, client):
        with app.app_context():
            make_user(username='phys', password='pw', roles='physicist')
        login(client, 'phys', 'pw')

        csv_file = io.BytesIO(b'equipment_class,facility,facility_address,eq_mod\nCT,New Site,,X\n')
        client.post('/import-data', data={'file': (csv_file, 'eq.csv')},
                    content_type='multipart/form-data')

        fac = Facility.query.filter_by(name='New Site').one()
        assert fac.address == ''
        assert fac.facility_full == ''

    def test_facility_import_without_address_column_keeps_address(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            db.session.add(Facility(name='St. Mary', address='1 Main St'))
            db.session.commit()
        login(client, 'admin', 'pw')

        csv_file = (io.BytesIO(b'name,facility_full\nSt. Mary,St. Mary Regional\n'), 'facilities.csv')
        client.post('/import-facilities', data={'csv_file': csv_file},
                    content_type='multipart/form-data')

        fac = Facility.query.filter_by(name='St. Mary').one()
        assert fac.address == '1 Main St'
        assert fac.facility_full == 'St. Mary Regional'


# ---------------------------------------------------------------------------
# Security: personnel permissions, CSRF, output escaping
# ---------------------------------------------------------------------------

def _personnel_form(person, **overrides):
    data = {'name': person.name, 'email': person.email, 'roles': person.get_roles_list() or ['contact']}
    data.update(overrides)
    return data


class TestPersonnelPermissions:
    def test_physicist_cannot_grant_self_admin(self, app, client):
        with app.app_context():
            me = make_user(username='phys', password='pw', roles='physicist')
            my_id = me.id
        login(client, 'phys', 'pw')
        client.post(f'/personnel/{my_id}/edit',
                    data=_personnel_form(me, roles=['physicist', 'admin'], is_admin='y', is_active='y',
                                         login_required='y', username='phys'))
        me = db.session.get(Personnel, my_id)
        assert me.is_admin is False
        assert 'admin' not in me.get_roles_list()

    def test_physicist_edit_leaves_login_fields_alone(self, app, client):
        with app.app_context():
            make_user(username='phys', password='pw', roles='physicist')
            other = make_user(username='other', password='original', roles='physics_assistant')
            other_id = other.id
        login(client, 'phys', 'pw')
        client.post(f'/personnel/{other_id}/edit',
                    data=_personnel_form(other, phone='555-1234', username='hijacked', password='newpassword1'))
        other = db.session.get(Personnel, other_id)
        assert other.phone == '555-1234'
        assert other.username == 'other'
        assert other.check_password('original')
        assert other.is_active and other.login_required

    def test_physicist_cannot_edit_or_delete_admin(self, app, client):
        with app.app_context():
            make_user(username='phys', password='pw', roles='physicist')
            admin = make_user(username='boss', password='pw', is_admin=True)
            admin_id = admin.id
        login(client, 'phys', 'pw')
        client.post(f'/personnel/{admin_id}/edit', data=_personnel_form(admin, name='Changed'))
        client.post(f'/personnel/{admin_id}/delete')
        admin = db.session.get(Personnel, admin_id)
        assert admin is not None and admin.name == 'Test User'

    def test_personnel_import_never_grants_admin(self, app, client):
        with app.app_context():
            me = make_user(username='pa', password='pw', roles='physics_assistant')
            my_id = me.id
        login(client, 'pa', 'pw')
        csv_file = io.BytesIO(f'id,name,email,roles\n{my_id},PA,pa@example.com,"physics_assistant, admin"\n'.encode())
        client.post('/import-personnel', data={'csv_file': (csv_file, 'p.csv')},
                    content_type='multipart/form-data')
        me = db.session.get(Personnel, my_id)
        assert me.is_admin is False
        assert 'admin' not in me.get_roles_list()

    def test_import_does_not_reactivate_existing_records(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            gone = make_user(username='gone', password='pw', roles='contact', is_active=False)
            gone_id = gone.id
        login(client, 'admin', 'pw')
        csv_file = io.BytesIO(b'name,email,roles\nGone,gone@example.com,contact\n')
        client.post('/import-personnel', data={'csv_file': (csv_file, 'p.csv')},
                    content_type='multipart/form-data')
        assert db.session.get(Personnel, gone_id).is_active is False

    def test_admin_cannot_demote_or_delete_self(self, app, client):
        with app.app_context():
            admin = make_user(username='admin', password='pw', is_admin=True, roles='physicist')
            admin_id = admin.id
        login(client, 'admin', 'pw')
        client.post(f'/personnel/{admin_id}/edit',
                    data=_personnel_form(admin, is_active='y', login_required='y', username='admin'))
        client.post(f'/personnel/{admin_id}/delete')
        admin = db.session.get(Personnel, admin_id)
        assert admin is not None and admin.is_admin is True

    def test_admin_can_grant_admin(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            other = make_user(username='other', password='pw', roles='physicist')
            other_id = other.id
        login(client, 'admin', 'pw')
        client.post(f'/personnel/{other_id}/edit',
                    data=_personnel_form(other, is_admin='y', is_active='y', login_required='y', username='other'))
        assert db.session.get(Personnel, other_id).is_admin is True


class TestActiveAndLoginAccess:
    def _deactivate(self, user_id, **changes):
        person = db.session.get(Personnel, user_id)
        for attr, value in changes.items():
            setattr(person, attr, value)
        db.session.commit()
        g.pop('_login_user', None)  # the test app context outlives requests; a real request reloads the user

    @pytest.mark.parametrize('changes', [{'is_active': False}, {'login_required': False}])
    def test_open_session_ends_when_access_is_removed(self, app, client, changes):
        with app.app_context():
            user_id = make_user(username='p', password='pw', roles='physicist').id
        login(client, 'p', 'pw')
        assert client.get('/personnel').status_code == 200
        self._deactivate(user_id, **changes)
        resp = client.get('/personnel')
        assert resp.status_code == 302 and '/login' in resp.location

    def test_login_access_needs_username(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            other = Personnel(name='New, Person', email='new@example.com', roles='contact')
            db.session.add(other)
            db.session.commit()
            other_id = other.id
        login(client, 'admin', 'pw')
        page = client.post(f'/personnel/{other_id}/edit',
                           data=_personnel_form(other, is_active='y', login_required='y')).data
        assert b'A username is required for login access.' in page
        assert db.session.get(Personnel, other_id).login_required is False

    def test_short_password_rejected(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            other = Personnel(name='New, Person', email='new@example.com', roles='contact')
            db.session.add(other)
            db.session.commit()
            other_id = other.id
        login(client, 'admin', 'pw')
        client.post(f'/personnel/{other_id}/edit',
                    data=_personnel_form(other, is_active='y', login_required='y', username='new', password='short'))
        assert db.session.get(Personnel, other_id).password_hash is None
        client.post(f'/personnel/{other_id}/edit',
                    data=_personnel_form(other, is_active='y', login_required='y', username='new',
                                         password='long enough pw'))
        assert db.session.get(Personnel, other_id).can_log_in

    def test_admin_cannot_remove_own_login_access(self, app, client):
        with app.app_context():
            admin = make_user(username='admin', password='pw', is_admin=True, roles='physicist')
            admin_id = admin.id
        login(client, 'admin', 'pw')
        client.post(f'/personnel/{admin_id}/edit',
                    data=_personnel_form(admin, is_admin='y', is_active='y', username='admin'))
        assert db.session.get(Personnel, admin_id).login_required is True

    def test_inactive_people_only_listed_when_assigned(self, app, client):
        with app.app_context():
            make_user(username='p', password='pw', roles='physicist')
            kept = make_user(username='kept', roles='contact', is_active=False)
            kept.name = 'Kept, Contact'
            gone = make_user(username='gone', roles='contact', is_active=False)
            gone.name = 'Gone, Contact'
            _seed_equipment()
            db.session.get(Equipment, 1).contact_id = kept.id
            db.session.commit()
        login(client, 'p', 'pw')
        names = [c['name'] for c in client.get('/api/equipment/1/form-data').get_json()['choices']['contacts']]
        assert names == ['Kept, Contact (inactive)']
        assert client.get('/api/equipment/2/form-data').get_json()['choices']['contacts'] == []

        # Saving the edit form keeps the inactive contact rather than rejecting or clearing it
        page = client.get('/equipment/1/edit').data.decode()
        assert 'Kept, Contact (inactive)' in page and 'Gone, Contact' not in page
        eq = db.session.get(Equipment, 1)
        client.post('/equipment/1/edit', data={'class_id': str(eq.class_id), 'contact_id': str(eq.contact_id),
                                               'eq_mod': 'Changed', 'eq_captype': 'Replacement'})
        eq = db.session.get(Equipment, 1)
        assert eq.eq_mod == 'Changed' and eq.contact_id is not None

    def test_inactive_reviewer_kept_on_existing_test(self, app, client):
        with app.app_context():
            make_user(username='p', password='pw', roles='physicist')
            reviewer = make_user(username='r', roles='physicist', is_active=False)
            reviewer.name = 'Former, Reviewer'
            _seed_equipment()
            db.session.add(ComplianceTest(eq_id=1, test_type='Annual', test_date=date(2024, 1, 1),
                                          reviewed_by_id=reviewer.id))
            db.session.commit()
        login(client, 'p', 'pw')
        assert 'Former, Reviewer (inactive)' in client.get('/compliance/test/1/edit').data.decode()
        assert 'Former, Reviewer' not in client.get('/compliance/test/1/new').data.decode()


class TestDatabaseErrorText:
    def test_import_reports_duplicate_without_raw_error(self, app, client):
        with app.app_context():
            admin = make_user(username='admin', password='pw', is_admin=True)
            make_user(username='taken', roles='contact')
            admin_id = admin.id
        login(client, 'admin', 'pw')
        csv_file = io.BytesIO(f'id,name,email\n{admin_id},Admin,taken@example.com\n'.encode())
        page = client.post('/import-personnel', data={'csv_file': (csv_file, 'p.csv')},
                           content_type='multipart/form-data', follow_redirects=True).data.decode()
        assert 'Row 2: another record already has this email' in page
        assert 'constraint' not in page.lower()

class TestCSRF:
    @pytest.fixture()
    def csrf_client(self, app):
        app.config['WTF_CSRF_ENABLED'] = True
        yield app.test_client()
        app.config['WTF_CSRF_ENABLED'] = False

    def _login(self, client):
        page = client.get('/login').data.decode()
        token = re.search(r'name="csrf-token" content="([^"]+)"', page).group(1)
        client.post('/login', data={'username': 'admin', 'password': 'pw', 'csrf_token': token})
        return token

    def test_post_without_token_is_rejected(self, app, csrf_client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            db.session.add(Facility(name='Keep'))
            db.session.commit()
        self._login(csrf_client)
        csrf_client.post('/admin/facilities/1/delete')
        assert db.session.get(Facility, 1).is_active is True

    def test_post_with_token_is_accepted(self, app, csrf_client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            db.session.add(Facility(name='Go'))
            db.session.commit()
        token = self._login(csrf_client)
        csrf_client.post('/admin/facilities/1/delete', data={'csrf_token': token})
        assert db.session.get(Facility, 1).is_active is False

    def test_json_post_without_token_gets_json_error(self, app, csrf_client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            cls = EquipmentClass(name='CT')
            db.session.add(cls)
            db.session.commit()
            db.session.add(Equipment(class_id=cls.id, eq_mod='Original'))
            db.session.commit()
        self._login(csrf_client)
        resp = csrf_client.post('/api/equipment/1/update-capital', json={'eq_capcst': '5'})
        assert resp.status_code == 400
        assert resp.get_json()['success'] is False


class TestOutputEscaping:
    def test_bubble_data_is_structured_and_survives_commas(self, app, client):
        with app.app_context():
            make_user(username='viewer', password='pw')
            cls = EquipmentClass(name='CT')
            fac = Facility(name='North, Wing')
            db.session.add_all([cls, fac])
            db.session.commit()
            db.session.add(Equipment(class_id=cls.id, facility_id=fac.id, eq_rm='1, 2',
                                     eq_eoldate=date(2000, 1, 1), eq_capcst=500, eq_radcap=1))
            db.session.commit()
        login(client, 'viewer', 'pw')
        points = client.get('/capital/bubble-data').get_json()['data']
        assert points[0]['facility'] == 'North, Wing'
        assert points[0]['room'] == '1, 2'
        assert points[0]['year'] == date.today().year
        assert points[0]['isEstimated'] is False

    def test_delete_button_keeps_test_type_out_of_js(self, app, client):
        payload = "x');alert(1);//"
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            cls = EquipmentClass(name='CT')
            db.session.add(cls)
            db.session.commit()
            eq = Equipment(class_id=cls.id)
            db.session.add(eq)
            db.session.commit()
            db.session.add(ComplianceTest(eq_id=eq.eq_id, test_type=payload, test_date=date(2024, 1, 1)))
            db.session.commit()
        login(client, 'admin', 'pw')
        page = client.get('/equipment/1').data.decode()
        handler = page.split('onclick="deleteComplianceTest(')[1].split('"')[0]
        assert handler == 'this.dataset.testId, this.dataset.testType, this.dataset.testDate)'
        assert 'data-test-type="x&#39;);alert(1);//"' in page

    def test_workbook_writes_formula_text_literally(self, app, client):
        formula = '=HYPERLINK("http://example.com","x")'
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            cls = EquipmentClass(name='CT')
            db.session.add(cls)
            db.session.commit()
            db.session.add(Equipment(class_id=cls.id, eq_notes=formula))
            db.session.commit()
        login(client, 'admin', 'pw')
        resp = client.get('/admin/database/workbook')
        workbook = openpyxl.load_workbook(io.BytesIO(resp.data))
        try:
            sheet = workbook['equipment']
            header = [c.value for c in sheet[1]]
            cell = sheet.cell(row=2, column=header.index('eq_notes') + 1)
            assert cell.data_type == 's'
            assert cell.value == formula
        finally:
            workbook.close()


# ---------------------------------------------------------------------------
# Filters, parameters, and inline edits
# ---------------------------------------------------------------------------

def _seed_equipment():
    """Two classes; CT at 'North' is active, planned, radiology-owned; MR is retired."""
    ct, mr = EquipmentClass(name='CT'), EquipmentClass(name='MR')
    north = Facility(name='North')
    db.session.add_all([ct, mr, north])
    db.session.commit()
    db.session.add_all([
        Equipment(class_id=ct.id, facility_id=north.id, eq_mod='Alpha', eq_planned=True, eq_radcap=1),
        Equipment(class_id=ct.id, eq_mod='Beta'),
        Equipment(class_id=mr.id, eq_mod='Gamma', eq_retired=True),
    ])
    db.session.commit()


def _exported_models(resp):
    rows = list(csv.DictReader(io.StringIO(resp.data.decode())))
    return sorted(r['eq_mod'] for r in rows)


class TestSharedFilters:
    def test_compliance_filter_plus_search_does_not_crash(self, app, client):
        with app.app_context():
            make_user(username='v', password='pw')
            _seed_equipment()
        login(client, 'v', 'pw')
        assert client.get('/compliance?eq_class=CT&eq_fac=North&search=Alpha').status_code == 200
        assert client.get('/export-compliance?eq_class=CT&eq_fac=North&search=Alpha').status_code == 200

    def test_export_matches_list_defaults(self, app, client):
        with app.app_context():
            make_user(username='v', password='pw')
            _seed_equipment()
        login(client, 'v', 'pw')
        # List view: planned and retired excluded unless asked for
        assert _exported_models(client.get('/export-equipment')) == ['Beta']
        assert _exported_models(client.get('/export-equipment?include_planned=true&include_retired=true')) == \
            ['Alpha', 'Beta', 'Gamma']
        assert _exported_models(client.get('/export-equipment?eq_class=CT&include_planned=true')) == ['Alpha', 'Beta']

    def test_capital_export_uses_capital_defaults(self, app, client):
        with app.app_context():
            make_user(username='v', password='pw')
            _seed_equipment()
        login(client, 'v', 'pw')
        # Capital view: planned included, radiology-owned only
        assert _exported_models(client.get('/export-equipment?view=capital')) == ['Alpha']

    def test_compliance_export_keeps_retired_history(self, app, client):
        with app.app_context():
            make_user(username='v', password='pw')
            _seed_equipment()
            db.session.add(ComplianceTest(eq_id=3, test_type='Annual', test_date=date(2020, 1, 1)))
            db.session.commit()
        login(client, 'v', 'pw')
        rows = list(csv.DictReader(io.StringIO(client.get('/export-compliance').data.decode())))
        assert [r['eq_id'] for r in rows] == ['3']

    def test_search_matches_full_facility_name(self, app, client):
        with app.app_context():
            make_user(username='v', password='pw')
            _seed_equipment()
            Facility.query.filter_by(name='North').one().facility_full = 'Northern Regional Hospital'
            db.session.commit()
        login(client, 'v', 'pw')
        resp = client.get('/export-equipment?search=Regional&include_planned=true')
        assert _exported_models(resp) == ['Alpha']


class TestBadParameters:
    def test_unknown_sort_field_and_bad_paging_do_not_crash(self, app, client):
        with app.app_context():
            make_user(username='v', password='pw')
            _seed_equipment()
        login(client, 'v', 'pw')
        for url in ('/equipment?sort=to_dict', '/equipment?sort=query,eq_id',
                    '/equipment?sort=days_until_due&page=-3&per_page=0',
                    '/capital?page=-1&per_page=-5', '/api/subclasses?class_id=abc'):
            assert client.get(url).status_code == 200, url

    def test_eq_id_in_query_string_does_not_break_redirect(self, app, client):
        with app.app_context():
            make_user(username='p', password='pw', roles='physicist')
            _seed_equipment()
        login(client, 'p', 'pw')
        resp = client.post('/schedule/test/2/new?eq_id=999&search=x',
                           data={'scheduled_date': '2030-01-01', 'scheduling_date': '2029-12-01'})
        assert resp.status_code == 302
        assert '/equipment/2' in resp.location and 'search=x' in resp.location


class TestEquipmentForms:
    def test_create_and_edit_with_dropdowns(self, app, client):
        with app.app_context():
            make_user(username='p', password='pw', roles='physicist')
            contact = make_user(username='c', password='pw', roles='contact')
            contact_id = contact.id
            _seed_equipment()
        login(client, 'p', 'pw')
        page = client.get('/equipment/new').data.decode()
        assert '<option value="1">CT</option>' in page and 'Select Contact' in page

        resp = client.post('/equipment/new', data={'class_id': '2', 'facility_id': '1', 'contact_id': str(contact_id),
                                                   'eq_mod': 'Delta', 'eq_captype': 'Replacement'})
        assert resp.status_code == 302
        eq = Equipment.query.filter_by(eq_mod='Delta').one()
        assert (eq.class_id, eq.facility_id, eq.contact_id, eq.manufacturer_id) == (2, 1, contact_id, None)

        edit_page = client.get(f'/equipment/{eq.eq_id}/edit').data.decode()
        assert '<option selected value="1">North</option>' in edit_page
        client.post(f'/equipment/{eq.eq_id}/edit', data={'class_id': '2', 'facility_id': '', 'eq_mod': 'Delta',
                                                          'eq_captype': 'Replacement'})
        assert db.session.get(Equipment, eq.eq_id).facility_id is None

    def test_inline_form_data_lists_choices(self, app, client):
        with app.app_context():
            make_user(username='p', password='pw', roles='physicist')
            _seed_equipment()
        login(client, 'p', 'pw')
        data = client.get('/api/equipment/1/form-data').get_json()
        assert [c['name'] for c in data['choices']['classes']] == ['CT', 'MR']
        assert data['choices']['facilities'] == [{'id': 1, 'name': 'North'}]
        assert 'Annual - TJC' in data['choices']['audit_frequencies']


class TestComplianceAndScheduleForms:
    def _login(self, app, client):
        with app.app_context():
            user = make_user(username='p', password='pw', roles='physicist')
            user.name = 'Bevins, Nick'
            db.session.commit()
            _seed_equipment()
            return user.id

    def test_compliance_test_add_edit_delete(self, app, client):
        user_id = self._login(app, client)
        login(client, 'p', 'pw')
        assert client.get('/compliance/test/2/new').status_code == 200
        resp = client.post('/compliance/test/2/new?search=x',
                           data={'test_type': 'Annual', 'test_date': '2024-05-01', 'reviewed_by_id': str(user_id)})
        assert '/equipment/2' in resp.location and 'search=x' in resp.location
        test = ComplianceTest.query.one()
        assert (test.eq_id, test.created_by, test.reviewed_by_id, test.performed_by_id) == (2, 'NB', user_id, None)

        resp = client.post(f'/compliance/test/{test.test_id}/edit?redirect_to=compliance',
                           data={'test_type': 'Audit', 'test_date': '2024-05-02'})
        assert resp.location.endswith('/compliance')
        assert db.session.get(ComplianceTest, test.test_id).test_type == 'Audit'

        resp = client.post(f'/compliance/test/{test.test_id}/delete', data={'search': 'y'})
        assert '/equipment/2' in resp.location and 'search=y' in resp.location
        assert ComplianceTest.query.count() == 0

    def test_scheduled_test_add_edit_delete(self, app, client):
        from app import ScheduledTest
        user_id = self._login(app, client)
        login(client, 'p', 'pw')
        client.post('/schedule/test/2/new', data={'scheduled_date': '2030-01-01', 'scheduling_date': '2029-12-01'})
        sched = ScheduledTest.query.one()
        assert (sched.eq_id, sched.created_by_id, sched.modified_by_id) == (2, user_id, user_id)
        client.post(f'/schedule/test/{sched.schedule_id}/edit',
                    data={'scheduled_date': '2030-02-01', 'scheduling_date': '2029-12-01'})
        assert db.session.get(ScheduledTest, sched.schedule_id).scheduled_date == date(2030, 2, 1)
        resp = client.post(f'/schedule/test/{sched.schedule_id}/delete', data={'redirect_to': 'compliance'})
        assert resp.location.endswith('/compliance')
        assert ScheduledTest.query.count() == 0


class TestInlineEdits:
    def _login(self, app, client):
        with app.app_context():
            make_user(username='p', password='pw', roles='physicist')
            _seed_equipment()
        login(client, 'p', 'pw')

    def test_retiring_inline_sets_retirement_date(self, app, client):
        self._login(app, client)
        resp = client.post('/api/equipment/2/update-details', json={'class_id': '1', 'eq_retired': True})
        assert resp.get_json()['success'] is True
        assert db.session.get(Equipment, 2).eq_retdate == date.today()

    def test_bad_values_are_rejected_or_ignored_not_500(self, app, client):
        self._login(app, client)
        resp = client.post('/api/equipment/2/update-details', json={'class_id': ''})
        assert resp.status_code == 400
        resp = client.post('/api/equipment/2/update-capital', json={'eq_capyr': '99999'})
        assert resp.status_code == 400
        resp = client.post('/api/equipment/2/update-capital', json={'eq_capcst': 'lots', 'eq_captype': 'Bogus'})
        assert resp.status_code == 200
        eq = db.session.get(Equipment, 2)
        assert eq.eq_capcst is None and eq.eq_captype == 'Replacement'
        resp = client.post('/api/equipment/2/update-contacts', json={'contact_id': 'x'})
        assert resp.status_code == 200

    def test_inline_audit_frequency_is_validated(self, app, client):
        self._login(app, client)
        client.post('/api/equipment/2/update-details',
                    json={'class_id': '1', 'eq_auditfreq': 'Quarterly, Whenever, Annual - ACR'})
        assert db.session.get(Equipment, 2).eq_auditfreq == 'Quarterly, Annual - ACR'


class TestCapitalCategoryOverlap:
    def test_range_touching_open_ended_category_is_rejected(self, app, client):
        from app import CapitalCategory
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            db.session.add(CapitalCategory(name='Big', min_cost=500, max_cost=None))
            db.session.commit()
        login(client, 'admin', 'pw')
        resp = client.post('/admin/capital-categories/add',
                           data={'name': 'Mid', 'min_cost': '100', 'max_cost': '500'}, follow_redirects=True)
        assert b'overlaps' in resp.data
        assert CapitalCategory.query.count() == 1


class TestAdminLookupTables:
    """The shared list/add/edit/deactivate/activate routes, for every lookup table."""

    # slug, model name, extra form fields for a valid add
    TABLES = [
        ('equipment-classes', 'EquipmentClass', {}),
        ('equipment-subclasses', 'EquipmentSubclass', {'class_id': '1', 'expected_lifetime': '10'}),
        ('departments', 'Department', {}),
        ('facilities', 'Facility', {'facility_full': 'Full Name', 'address': '1 Main St'}),
        ('manufacturers', 'Manufacturer', {}),
        ('capital-categories', 'CapitalCategory', {'min_cost': '0', 'max_cost': '99'}),
    ]

    def _login(self, app, client):
        with app.app_context():
            make_user(username='admin', password='pw', is_admin=True)
            db.session.add(EquipmentClass(name='Base'))  # id 1, parent for subclasses
            db.session.commit()
        login(client, 'admin', 'pw')

    @pytest.mark.parametrize('slug, model_name, extra', TABLES)
    def test_full_lifecycle(self, app, client, slug, model_name, extra):
        import app as app_module
        model = getattr(app_module, model_name)
        self._login(app, client)
        base = f'/admin/{slug}'
        assert client.get(base).status_code == 200
        assert client.get(f'{base}/add').status_code == 200

        client.post(f'{base}/add', data={'name': 'Alpha', **extra})
        item = model.query.filter_by(name='Alpha').one()
        for attr, value in extra.items():
            assert str(getattr(item, attr)) == value

        page = client.get(f'{base}/{item.id}/edit').data.decode()
        assert 'value="Alpha"' in page and 'None' not in page
        client.post(f'{base}/{item.id}/edit', data={'name': 'Beta', **extra})
        assert db.session.get(model, item.id).name == 'Beta'

        client.post(f'{base}/{item.id}/delete')
        assert db.session.get(model, item.id).is_active is False
        assert b'Inactive' in client.get(base).data
        client.post(f'{base}/{item.id}/activate')
        assert db.session.get(model, item.id).is_active is True

    def test_duplicate_name_is_rejected_and_input_kept(self, app, client):
        self._login(app, client)
        client.post('/admin/departments/add', data={'name': 'Radiology'})
        resp = client.post('/admin/departments/add', data={'name': 'Radiology'})
        assert b'already exists' in resp.data and b'value="Radiology"' in resp.data
        assert Department.query.count() == 1

    def test_readding_a_deactivated_name_reactivates_it(self, app, client):
        self._login(app, client)
        client.post('/admin/facilities/add', data={'name': 'North'})
        fac = Facility.query.one()
        client.post(f'/admin/facilities/{fac.id}/delete')
        resp = client.post('/admin/facilities/add', data={'name': 'North', 'address': 'New Address'},
                           follow_redirects=True)
        assert b'reactivated' in resp.data
        fac = Facility.query.one()
        assert fac.is_active and fac.address == 'New Address'

    def test_subclass_names_are_unique_only_within_a_class(self, app, client):
        from app import EquipmentSubclass
        self._login(app, client)
        with app.app_context():
            db.session.add(EquipmentClass(name='Other'))  # id 2
            db.session.commit()
        client.post('/admin/equipment-subclasses/add', data={'name': 'Head', 'class_id': '1'})
        client.post('/admin/equipment-subclasses/add', data={'name': 'Head', 'class_id': '2'})
        resp = client.post('/admin/equipment-subclasses/add', data={'name': 'Head', 'class_id': '1'})
        assert b'already exists' in resp.data
        assert EquipmentSubclass.query.count() == 2

    def test_invalid_number_and_missing_required_field(self, app, client):
        from app import CapitalCategory, EquipmentSubclass
        self._login(app, client)
        resp = client.post('/admin/capital-categories/add', data={'name': 'A', 'min_cost': 'lots'})
        assert b'must be a whole number' in resp.data
        resp = client.post('/admin/equipment-subclasses/add', data={'name': 'Head', 'class_id': ''})
        assert b'Equipment Class is required' in resp.data
        assert CapitalCategory.query.count() == 0 and EquipmentSubclass.query.count() == 0

    def test_capital_max_must_exceed_min(self, app, client):
        from app import CapitalCategory
        self._login(app, client)
        resp = client.post('/admin/capital-categories/add', data={'name': 'A', 'min_cost': '100', 'max_cost': '50'})
        assert b'Maximum cost must be greater' in resp.data
        assert CapitalCategory.query.count() == 0

    def test_edit_with_no_changes(self, app, client):
        self._login(app, client)
        resp = client.post('/admin/equipment-classes/1/edit', data={'name': 'Base'}, follow_redirects=True)
        assert b'No changes made' in resp.data

    def test_non_admin_is_redirected(self, app, client):
        with app.app_context():
            make_user(username='phys', password='pw', roles='physicist')
            db.session.add(Department(name='Keep'))
            db.session.commit()
        login(client, 'phys', 'pw')
        assert client.get('/admin/departments').status_code == 302
        client.post('/admin/departments/1/delete')
        assert db.session.get(Department, 1).is_active is True
